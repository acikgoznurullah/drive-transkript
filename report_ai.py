"""
report_ai.py – Transkriptten, kullanıcının örnek raporlarına uygun Word raporu yazar (Claude API).

Claude, RAPOR_FORMATLARI.md kurallarına ve ornekler/ klasöründeki örnek raporlara bakarak report_builder.py'nin
JSON şemasında içerik üretir; tasarım report_builder.py'de sabittir.

Kullanım (elle deneme):
    python report_ai.py <transkript.txt> <cikti.docx> [kategori] [dil]
Gerekli ortam değişkeni: ANTHROPIC_API_KEY
"""
from __future__ import annotations
import json, re, sys
from datetime import datetime
from pathlib import Path

import anthropic
import docx

import report_builder

HERE = Path(__file__).resolve().parent
GUIDE = HERE / "RAPOR_FORMATLARI.md"
SAMPLES = HERE / "ornekler"
MODEL = "claude-opus-5-5"


def docx_text(path: Path) -> str:
    """Örnek raporun stil adlarıyla birlikte düz metni (Heading 1 | Başlık …), tablolar dahil."""
    d = docx.Document(str(path))
    lines = [f"{p.style.name if p.style is not None else '-'} | {p.text}" for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append("TABLO | " + " ‖ ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def system_prompt() -> str:
    schema = report_builder.__doc__.split("JSON şeması", 1)[1]
    samples = "\n\n".join(f'<ornek dosya="{p.name}">\n{docx_text(p)}\n</ornek>'
                          for p in sorted(SAMPLES.glob("*.docx")))
    return f"""Sen, ses kayıtlarının otomatik dökümlerinden kullanıcının örnek raporlarıyla aynı üslupta
düzenlenmiş Word raporları hazırlayan bir editörsün. Görsel tasarım sabittir; senin işin yalnızca içerik
JSON'unu üretmektir.

<rapor_formatlari>
{GUIDE.read_text(encoding="utf-8")}
</rapor_formatlari>

<json_semasi>
JSON şeması{schema}
</json_semasi>

<ornek_raporlar>
Aşağıdakiler kullanıcının beğendiği örnek raporların metinleridir (stil adı | metin). Bölüm kurgusunu, başlık
üslubunu, ayrıntı düzeyini ve uzunluğu bunlara göre ayarla.
{samples}
</ornek_raporlar>

Çıktı kuralları:
- Yalnızca tek bir JSON nesnesi döndür, başka açıklama yazma. JSON'u ```json bloğu içine koy.
- Transkriptte olmayan bilgi, isim, rakam ve tarih ekleme. Dökümdeki bariz ses tanıma hatalarını düzelt ve
  bunu "note" alanında belirt.
- Rapor kaydın dilinde yazılır; dil "en" ise "lang": "en" ve İngilizce rapor."""


def date_from_name(source: str) -> str | None:
    m = re.search(r"(\d{2})(\d{2})(\d{2})_\d{6}", source)
    if not m:
        return None
    yy, mm, dd = (int(x) for x in m.groups())
    try:
        return datetime(2000 + yy, mm, dd).strftime("%Y-%m-%d")
    except ValueError:
        return None


def make_report(txt: Path, out_docx: Path, source: str, category: str | None = None,
                language: str | None = None, duration: str | None = None) -> Path:
    transcript = txt.read_text(encoding="utf-8")
    meta = {"dosya": source, "kategori": category, "dil": language or "tr", "süre": duration,
            "kayıt_tarihi (dosya adından)": date_from_name(source),
            "transkripsiyon_günü": datetime.now().strftime("%Y-%m-%d")}
    user = (f"Kayıt bilgileri:\n{json.dumps(meta, ensure_ascii=False, indent=2)}\n\n"
            f"Kategoriye ve içeriğe göre uygun rapor türünü seç ve raporun içerik JSON'unu yaz.\n\n"
            f"<transkript>\n{transcript}\n</transkript>")

    client = anthropic.Anthropic(max_retries=4)
    with client.beta.messages.stream(
        model=MODEL,
        max_tokens=64000,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        output_config={"effort": "high"},
        system=[{"type": "text", "text": system_prompt(), "cache_control": {"type": "ephemeral"}}],
        messages=[{"role": "user", "content": user}],
    ) as stream:
        msg = stream.get_final_message()

    if msg.stop_reason == "refusal":
        raise RuntimeError(f"Model raporu yazmayı reddetti: {getattr(msg.stop_details, 'category', None)}")
    if msg.stop_reason == "max_tokens":
        raise RuntimeError("Yanıt uzunluk sınırına takıldı.")
    text = "".join(b.text for b in msg.content if b.type == "text")
    m = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.S)
    spec = json.loads(m.group(1) if m else text[text.find("{"): text.rfind("}") + 1])

    spec_path = out_docx.with_suffix(".json")
    spec_path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    report_builder.build(spec, str(out_docx))
    if len(docx.Document(str(out_docx)).paragraphs) == 0:
        raise RuntimeError("Oluşan rapor boş.")
    spec_path.unlink(missing_ok=True)
    return out_docx


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    t = Path(sys.argv[1])
    make_report(t, Path(sys.argv[2]), source=t.name,
                category=sys.argv[3] if len(sys.argv) > 3 else None,
                language=sys.argv[4] if len(sys.argv) > 4 else None)
    print("OK", sys.argv[2])
