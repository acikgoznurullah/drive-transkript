"""
report_builder.py – JSON içerik dosyasından, kullanıcının örnek raporlarıyla aynı tasarımda
(Drive'daki İş/Raporlar ve Özel/Raporlar örnekleri: DASK toplantı raporu, tefsir dersi raporları)
Word raporu üretir.

Kullanım:
    fwenv\\Scripts\\python.exe automation\\report_builder.py <icerik.json> <cikti.docx>

Tasarım: A4, 2,5 cm kenar; ortalanmış lacivert başlık, mavi alt başlık, gri meta satırı ve altın çizgi;
lacivert ve altın alt çizgili H1, mavi H2; iki yana yaslı 11 pt Calibri gövde; lacivert başlıklı tablolar;
"<rapor adı> | Sayfa N" altbilgisi.

JSON şeması (tüm metin alanlarında **kalın** işaretlemesi kullanılabilir):
{
  "lang": "tr",                                   # "tr" veya "en" (altbilgi "Sayfa"/"Page")
  "title": "Toplantı Raporu",
  "subtitle": "DASK için Sermaye Tabanlı Prim Metodolojisi Sunumu",
  "meta": ["23 Temmuz 2026", "T-Rupt Technology → DASK"],     # " • " ile birleştirilir
  "footer": "DASK Prim Metodolojisi — Toplantı Raporu",
  "info": [["Tarih / süre", "..."], ["Konu", "..."]],           # isteğe bağlı bilgi tablosu
  "info_headers": ["Alan", "Bilgi"],                           # isteğe bağlı
  "info_note": "İsimler kayıttan aktarılmıştır; ...",          # tablonun altında küçük gri not
  "intro": ["Başlıksız giriş paragrafı ..."],                  # isteğe bağlı
  "sections": [
    {"heading": "1. Toplantının Amacı",
     "paragraphs": ["..."],
     "bullets": ["düz madde", {"bold": "Etiket: ", "text": "açıklama"}],
     "numbered": ["numaralı madde", {"bold": "Sorumlu — ", "text": "görev"}],
     "table": {"headers": ["#", "Aksiyon", "Sorumlu"], "rows": [["1", "...", "..."]],
               "widths": [1.0, 10.5, 4.5]},            # cm; toplam ≈ 16 cm
     "after": ["bloktan sonra gelen paragraf"],
     "subsections": [ {...aynı alanlar...} ]}
  ],
  "note": "Not: Bu rapor ... dökümünden düzenlenmiştir."
}
Bir bölümde alanlar şu sırayla basılır: heading, paragraphs, bullets, numbered, table, after, subsections.
"""
from __future__ import annotations
import copy, json, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, BLUE, GOLD, GRAY, ROW = "1F3864", "2E5496", "C9A227", "808080", "F2F5FA"
FONT = "Calibri"
CONTENT_CM = 16.0          # A4 21 cm - 2 × 2,5 cm


# ------------------------------------------------------------------ düşük seviye yardımcılar
def rgb(hex6):
    return RGBColor.from_string(hex6)


def set_font(rpr_owner, name=FONT):
    """Tema fontlarını temizleyip açık font ataması yapar (Word aksi halde tema fontunu kullanır)."""
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), name)


def border(p, side, color, sz, space):
    pPr = p._p.get_or_add_pPr() if hasattr(p, "_p") else p
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr"); pPr.append(bdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    bdr.append(el)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)


def table_borders(table, color="BFBFBF", sz=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)


_BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_rich(p, text, bold=False, size=None, color=None, italic=False):
    """'**kalın**' işaretlemesini destekleyerek metin ekler."""
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            _run(p, text[pos:m.start()], bold, size, color, italic)
        _run(p, m.group(1), True, size, color, italic)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:], bold, size, color, italic)


def _run(p, text, bold, size, color, italic):
    r = p.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = rgb(color)
    return r


# ------------------------------------------------------------------ stiller
def setup_styles(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    set_font(normal.element)
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_after = Pt(6); pf.line_spacing = 1.15

    specs = {"Heading 1": (13.5, NAVY, 14, 6.5), "Heading 2": (11.5, BLUE, 9, 4.5)}
    for name, (size, color, before, after) in specs.items():
        st = doc.styles[name]
        set_font(st.element)
        st.font.size = Pt(size); st.font.bold = True; st.font.italic = False
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            border(st.element.get_or_add_pPr(), "bottom", GOLD, 6, 4)

    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def restart_numbering(doc, paragraph):
    """Her numaralı liste 1'den başlasın diye 'List Number' için yeni bir num örneği oluşturur."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_numpr = doc.styles["List Number"].element.pPr.numPr
    base_num_id = style_numpr.numId.val
    base = numbering.num_having_numId(base_num_id)
    abstract_id = base.abstractNumId.val
    num = numbering.add_num(abstract_id)
    ov = num.add_lvlOverride(ilvl=0)
    ov.add_startOverride(1)
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = num.numId
    return num.numId


def use_numbering(paragraph, num_id):
    numPr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = num_id


# ------------------------------------------------------------------ bloklar
def body_par(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_rich(p, text)
    return p


def item_text(p, it):
    if isinstance(it, dict):
        add_rich(p, it.get("bold", ""), bold=True)
        add_rich(p, it.get("text", ""))
    else:
        add_rich(p, str(it))


def add_table(doc, spec):
    headers, rows = spec["headers"], spec.get("rows", [])
    n = len(headers)
    widths = spec.get("widths") or [CONTENT_CM / n] * n
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t); cell_margins(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        add_rich(c.paragraphs[0], str(h), bold=True, size=9.5, color="FFFFFF")
        shade(c, NAVY)
    repeat_header(t.rows[0])
    bold_cols = set(spec.get("bold_cols", []))
    for row in rows:
        cells = t.add_row().cells
        for i in range(n):
            cells[i].text = ""
            add_rich(cells[i].paragraphs[0], str(row[i]) if i < len(row) else "", bold=i in bold_cols, size=9.5)
            shade(cells[i], ROW)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_block(doc, block, level):
    if block.get("heading"):
        doc.add_heading(block["heading"], level=level)
    for para in block.get("paragraphs", []):
        body_par(doc, para)
    for it in block.get("bullets", []):
        p = doc.add_paragraph(style="List Bullet"); item_text(p, it)
    if block.get("numbered"):
        num_id = None
        for it in block["numbered"]:
            p = doc.add_paragraph(style="List Number")
            num_id = restart_numbering(doc, p) if num_id is None else (use_numbering(p, num_id) or num_id)
            item_text(p, it)
    if block.get("table"):
        add_table(doc, block["table"])
    for para in block.get("after", []):
        body_par(doc, para)
    for sub in block.get("subsections", []):
        add_block(doc, sub, min(level + 1, 2))


def add_footer(doc, text, lang):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.text = ""
    border(p, "top", GOLD, 4, 6)
    # Footer stilinin varsayılan orta/sağ sekme durakları (Letter ölçüsüne göre) "Sayfa N"yi ortaya iter
    footer_style = doc.styles["Footer"].element.pPr
    tabs = footer_style.find(qn("w:tabs")) if footer_style is not None else None
    if tabs is not None:
        footer_style.remove(tabs)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), WD_TAB_ALIGNMENT.RIGHT)
    _run(p, text, False, 8, GRAY, False)
    _run(p, "\t" + ("Page " if lang == "en" else "Sayfa "), False, 8, GRAY, False)
    r = _run(p, "", False, 8, GRAY, False)
    for kind, val in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = val
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = val
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


# ------------------------------------------------------------------ ana
def build(spec: dict, out_path: str):
    doc = Document()
    setup_styles(doc)
    lang = spec.get("lang", "tr")

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = _run(p, spec.get("title", "Rapor"), True, 19, NAVY, False); set_font(r._r)
    if spec.get("subtitle"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        add_rich(p, spec["subtitle"], size=12.5, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    if spec.get("meta"):
        add_rich(p, "  •  ".join(spec["meta"]), size=9, color=GRAY)
    border(p, "bottom", GOLD, 6, 6)

    if spec.get("info"):
        headers = spec.get("info_headers") or (["Field", "Details"] if lang == "en" else ["Alan", "Bilgi"])
        add_table(doc, {"headers": headers, "rows": spec["info"], "widths": [4.2, CONTENT_CM - 4.2],
                        "bold_cols": [0, 1]})
        if spec.get("info_note"):
            doc.paragraphs[-1]._p.getparent().remove(doc.paragraphs[-1]._p)   # tablo sonrası boşluğu not alsın
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
            add_rich(p, spec["info_note"], size=8, color=GRAY, italic=True)

    for para in spec.get("intro", []):
        body_par(doc, para)

    for block in spec.get("sections", []):
        add_block(doc, block, 1)

    if spec.get("note"):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
        add_rich(p, spec["note"], size=9, color=GRAY, italic=True)

    add_footer(doc, spec.get("footer") or spec.get("title", ""), lang)
    doc.core_properties.title = spec.get("title", "")
    doc.save(out_path)


if __name__ == "__main__":
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        sys.exit("Kullanım: report_builder.py <icerik.json> <cikti.docx>")
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
    print(f"OK {sys.argv[2]}")
